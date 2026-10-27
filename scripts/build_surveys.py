import openpyxl, json, re, pickle, datetime, unicodedata
from shapely.geometry import Point, shape
import build_geo_fix as BF
G=pickle.load(open('geo.pkl','rb')); districts=G['districts']; lake=G['lake']
b=json.load(open('basemap.json')); cauce=shape(b['cauceAcaray']['features'][0]['geometry'])
DNAMES={BF.FIX.get(n,n) for n,d,g in districts}
def nrm(s): return re.sub(r'[^a-z0-9]','',unicodedata.normalize('NFD',str(s)).encode('ascii','ignore').decode().lower())
ALIAS={nrm(x):x for x in DNAMES}
ALIAS[nrm('Mariscal Francisco Solano López')]='Mcal. Francisco S. López'
ALIAS[nrm('Dr. J. Eulogio Estigarribia')]='Dr. J. Eulogio Estigarribia'
ALIAS[nrm('Tembiaporá')]='Tembiaporã'
def dist_at(p):
    for n,d,g in districts:
        if g.contains(p): return BF.FIX.get(n,n), d
    return '',''
def clean(v):
    if v is None: return ''
    if isinstance(v,float):
        if v.is_integer(): return int(v)
        return v
    s=str(v).replace('\xa0',' ').strip()
    s=re.sub(r'\s+',' ',s)
    return s
def date_from(photo):
    m=re.match(r'(\d{13})',str(photo or ''))
    if not m: return ''
    return datetime.datetime.utcfromtimestamp(int(m.group(1))/1000).strftime('%Y-%m-%d')
def emb_of(p):
    dl=0 if lake.contains(p) else lake.exterior.distance(p)
    dc=0 if cauce.contains(p) else cauce.boundary.distance(p)
    return 'Yguazú' if dl<=dc else 'Acaray'
wb=openpyxl.load_workbook('enc/encuestas.xlsx',data_only=True)
out=[]
def rows(ws):
    hdr=[clean(c.value) for c in ws[1]]
    for r in ws.iter_rows(min_row=2,values_only=True):
        if not any(v not in (None,'') for v in r): continue
        yield {h:clean(v) for h,v in zip(hdr,r) if h}
def find(d,*keys):
    for k in d:
        for key in keys:
            if nrm(k).startswith(nrm(key)): return d[k]
    return ''
DROP=('id','entrevistado','nombredelentrevistado','telefono','ubicacion','registrofotografico','agregarregistro','primerregistro','segundoregistro','tercerregistro','6registro','departamento','distrito')
def extras(d,skip=()):
    e={}
    for k,v in d.items():
        nk=nrm(k)
        if any(nk.startswith(x) for x in DROP) or k in skip or v in ('',None): continue
        e[k]=v
    return e
def base(d,code,tema):
    lat=float(find(d,'Ubicación_latitude','Ubicación GPS_latitude','Ubicacion _latitude','Ubicacion_latitude')); lon=float(find(d,'Ubicación_longitude','Ubicación GPS_longitude','Ubicacion _longitude','Ubicacion_longitude'))
    p=Point(lon,lat); dd,dep=dist_at(p)
    rep=find(d,'Distrito'); dist=ALIAS.get(nrm(rep),'') or dd
    photo=find(d,'Agregar registro','Primer registro','Registro fotográfico','6. Registro')
    return {"codigo":code,"fecha_relevamiento":date_from(photo),"encuestador":"Equipo social ANDE","embalse":emb_of(p),
            "departamento":find(d,'Departamento') or dep,"distrito":dist,"latitud":round(lat,6),"longitud":round(lon,6),
            "categoria_sitio":"Encuesta de depleción (ago. 2025) · "+tema,"estado_ficha":"Visitada por el equipo"}
FUENTE={"pozosomeroaljibe":"Pozo somero propio","pozoartesianocomunitario":"Pozo profundo comunitario","pozoartesianoparticular":"Pozo profundo particular","aguateriacomunitariaconapoyodeessap":"Aguatería comunitaria"}
for d in rows(wb['Acceso al agua']):
    r=base(d,f"EN-AA-{int(d['ID']):03d}","Acceso al agua")
    r.update({"tipo_unidad":"Hogar","localidad":d.get('Comunidad',''),"fuente_agua_consumo":FUENTE.get(nrm(d.get('Tipo Fuente de Agua','')),d.get('Tipo Fuente de Agua','')),
              "observaciones":d.get('Observaciones','')})
    r["_extra"]=extras(d,skip=('Comunidad','Tipo Fuente de Agua','Observaciones'))
    out.append(r)
for d in rows(wb['Agricultura y Ganadería']):
    r=base(d,f"EN-AP-{int(d['ID']):03d}","Agricultura y ganadería")
    tec=nrm(d.get('Técnicas de cultivo','')); cult=nrm(d.get('Cultivos realizados',''))
    act="Producción mecanizada" if 'mecaniz' in tec else ("Agricultura familiar" if cult and cult!='no' else ("Ganadería menor" if nrm(d.get('Tipos de animales',''))not in('','ninguno') else "Otra"))
    uso=str(find(d,'Uso de agua del embalse')); nu=nrm(uso); usos=[]
    if not nu.startswith('no') and nu:
        if 'animal' in nu or 'ganad' in nu or 'vaca' in nu: usos.append("Abrevadero de ganado")
        if 'riego' in nu: usos.append("Riego")
    r.update({"tipo_unidad":"Unidad productiva","nombre_sitio":d.get('Nombre del establecimiento o finca',''),"actividad_principal":act,"usos_agua":"; ".join(usos),"observaciones":d.get('Observaciones','')})
    r["_extra"]=extras(d,skip=('Nombre del establecimiento o finca','Observaciones'))
    out.append(r)
for d in rows(wb['Pesca']):
    r=base(d,f"EN-PE-{int(d['ID']):03d}","Pesca")
    ta=nrm(d.get('Tipo de actividad','')); dest=nrm(d.get('Destino del pescado',''))
    dep={"ingresoprincipal":"Alta","ingresocomplementario":"Media","actividadrecreativa":"Baja"}.get(ta,"")
    usos=["Pesca comercial" if ('mercado' in dest or 'mayorista' in dest) else "Pesca de subsistencia"]
    if ta=="actividadrecreativa": usos.append("Recreación")
    r.update({"tipo_unidad":"Hogar","localidad":find(d,'Comunidad o zona'),"actividad_principal":"Pesca","dependencia_embalse":dep,"usos_agua":"; ".join(usos),"afectaciones_previstas":"Pesca","observaciones":d.get('Observaciones','')})
    r["_extra"]=extras(d,skip=('Comunidad o zona','Nombre del entrevistado o finca','Observaciones'))
    out.append(r)
for d in rows(wb['Otras Observaciones']):
    r=base(d,f"EN-OB-{int(d['ID']):03d}","Otras observaciones")
    tr=nrm(d.get('Tipo de reporte',''))+nrm(d.get('Descripción',''))
    if re.search('recrea|turis|balnear|playa|camping|veranieg|hospedaje',tr): tipo="Balneario o recreativo"
    elif re.search('cultiv|ganad|tambo|empresa|agro|silo|forest|eucal|granja|acuic|piscic|arener|cantera|industr|plantac|banana|producc',tr): tipo="Unidad productiva"
    elif re.search('indigen',tr): tipo="Comunidad indígena"
    else: tipo="Sitio urbano o civil"
    desc=str(d.get('Descripción',''))
    corto=desc if len(desc)<=70 else desc[:67].rsplit(' ',1)[0]+'…'
    r.update({"tipo_unidad":tipo,"nombre_sitio":corto,"observaciones":("Tipo de reporte: "+str(d.get('Tipo de reporte','')))+(". "+desc if len(desc)>70 else "")})
    r["categoria_sitio"]+=" · "+str(d.get('Tipo de reporte',''))
    r["_extra"]={}
    out.append(r)
for r in out:
    r['observaciones']=re.sub(r'Gustavito y Pirulo|Javier y Ramón','[nombres omitidos]',r.get('observaciones',''))
json.dump(out,open('surveys.json','w'),ensure_ascii=False,separators=(',',':'))
import collections
print(len(out), collections.Counter(r['tipo_unidad'] for r in out), collections.Counter(r['distrito'] for r in out), collections.Counter(r['embalse'] for r in out), collections.Counter(r['fecha_relevamiento'] for r in out).most_common(4))
print(json.dumps(out[0],ensure_ascii=False)[:900]); print(json.dumps(out[50],ensure_ascii=False)[:700])
print([ (r['nombre_sitio'],r['tipo_unidad']) for r in out if r['codigo'].startswith('EN-OB')])
