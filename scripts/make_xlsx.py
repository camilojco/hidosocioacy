import json, openpyxl
from openpyxl.styles import Font, PatternFill
e2=json.load(open('entrega2.json'))
# Encabezados "como los entrega la encuestadora" (exportación tipo Kobo), distintos a los del visor
HEAD=[("Código","codigo"),("Fecha","fecha_relevamiento"),("Encuestador","encuestador"),("Embalse","embalse"),("Departamento","departamento"),("Distrito","distrito"),("Compañía","localidad"),
("_gps_latitude","latitud"),("_gps_longitude","longitud"),("_gps_precision","precision_gps_m"),("Tipo de unidad","tipo_unidad"),("Nombre y apellido","nombre_referente"),("Jefatura","sexo_jefatura"),
("Total personas","integrantes"),("Menores de 14 años","menores_14"),("Mayores de 65 años","mayores_65"),("Tenencia del terreno","tenencia"),("Franja de dominio","franja_dominio"),("Distancia a la costa (m)","distancia_costa_m"),
("Fuente de agua","fuente_agua_consumo"),("Usos del agua","usos_agua"),("Actividad principal","actividad_principal"),("Dependencia económica","dependencia_embalse"),("Conoce el descenso","conoce_descenso"),
("Preocupación","preocupacion_descenso"),("Afectaciones","afectaciones_previstas"),("Reclamo previo","reclamo_previo"),("Estado","estado_ficha"),("Observaciones","observaciones")]
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Censo_entrega_2"
ws.append([h for h,_ in HEAD]+["_uuid","Material de vivienda"])
import uuid, random
random.seed(7)
for r in e2:
    row=[]
    for h,k in HEAD:
        v=r.get(k,"")
        if k=="conoce_descenso" or k=="reclamo_previo": v={"Sí":"si","No":"no"}.get(v,v)
        row.append(v)
    row+= [str(uuid.UUID(int=random.getrandbits(128))), random.choice(["Ladrillo","Madera","Mixta","Ladrillo"])]
    ws.append(row)
for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="DDE7E6")
ws.freeze_panes="A2"
for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=max(12,min(34,len(str(col[0].value))+4))
wb.save("entrega_2_encuestadora.xlsx"); print("ok", len(e2))
